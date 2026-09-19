"""Business-data adapter (Person B, part 1).

Loads the ONE useful part of the chaotic Excel — the supplier master, the 2026
purchase orders, and the payment rules text — into clean lookups. Ignores the
~11 junk/trap sheets. De-dups the duplicated P007 master row.

Everything the rules engine needs about "who is a real supplier / what was
ordered" comes from here; "was it already paid / still pending" comes from the
ERP snapshot (Person C).
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

from .models import PurchaseOrder, Supplier

DEFAULT_XLSX = Path(__file__).resolve().parents[1] / "challenge" / "FINAL_v7_DEFINITIVO_ahorasi.xlsx"

SHEET_SUPPLIERS = "Proveedores"
SHEET_ORDERS = "Pedidos_2026"
SHEET_RULES = "Norma_Pagos_v3"

# --- schema resilience -------------------------------------------------------
# The master Excel is human-maintained, so its LAYOUT can drift (a column gets
# renamed / reordered, a sheet gets a slightly different name). We resolve
# columns by HEADER NAME (accent/case/space-insensitive) instead of by fixed
# position, so re-ordering or renaming a column no longer silently reads the
# wrong data. If a *required* sheet or column can't be located we raise a clear,
# actionable BusinessDataError instead of guessing.


class BusinessDataError(ValueError):
    """The master Excel is structurally unreadable (missing sheet/column)."""


# sheet-name aliases (normalised); the preferred name is tried first.
SUPPLIER_SHEET_ALIASES = ["proveedores", "proveedor", "suppliers", "maestro_proveedores", "maestro"]
ORDER_SHEET_ALIASES = ["pedidos_2026", "pedidos", "orders", "po_2026", "pedidos2026"]

# field -> ordered header aliases (normalised). first match wins.
SUPPLIER_COLUMNS: dict[str, list[str]] = {
    "id": ["id", "id_proveedor", "codigo", "cod_proveedor", "proveedor_id", "proveedorid"],
    "name": ["razon_social", "nombre", "proveedor", "name"],
    "tax_id": ["nif", "cif", "nif_cif", "cif_nif", "tax_id", "identificacion_fiscal"],
    "iban": ["iban", "cuenta", "iban_cuenta", "cuenta_bancaria"],
    "payment_terms": ["condiciones", "condiciones_pago", "forma_pago", "terminos_pago", "payment_terms", "plazo", "plazo_pago"],
}
SUPPLIER_REQUIRED = ("id", "name", "tax_id", "iban")

ORDER_COLUMNS: dict[str, list[str]] = {
    "id": ["pedido", "id_pedido", "num_pedido", "numero_pedido", "orden", "po", "order_id"],
    "supplier_id": ["proveedorid", "proveedor_id", "id_proveedor", "proveedor", "supplier_id"],
    "tax_id": ["nif", "cif", "tax_id"],
    "amount": ["importe_total", "importe", "total", "amount", "monto", "importe_pedido"],
}
ORDER_REQUIRED = ("id", "supplier_id", "amount")


def _norm_header(value: object) -> str:
    """Fold a header cell to a comparable key: lowercase, strip accents, snake_case."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = text.strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _find_sheet(wb, preferred: str, aliases: list[str], warnings: list[str], label: str):
    """Locate a sheet by its preferred name, then by normalised aliases."""
    if preferred in wb.sheetnames:
        return wb[preferred]
    norm_to_actual = {_norm_header(name): name for name in wb.sheetnames}
    for alias in aliases:
        actual = norm_to_actual.get(alias)
        if actual:
            warnings.append(f"hoja de {label}: se esperaba '{preferred}', se uso '{actual}'")
            return wb[actual]
    raise BusinessDataError(
        f"no se encontro la hoja de {label} (esperaba '{preferred}' o similar). "
        f"hojas disponibles: {', '.join(wb.sheetnames)}"
    )


def _header_index(ws) -> dict[str, int]:
    """Map each normalised header in row 1 to its column position (first wins)."""
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        idx: dict[str, int] = {}
        for i, cell in enumerate(row):
            key = _norm_header(cell)
            if key and key not in idx:
                idx[key] = i
        return idx
    return {}


def _resolve_columns(
    header_idx: dict[str, int],
    columns: dict[str, list[str]],
    required: tuple[str, ...],
    label: str,
) -> dict[str, int]:
    """Resolve field -> column index via header aliases; error if a required one is missing."""
    resolved: dict[str, int] = {}
    for field_name, aliases in columns.items():
        for alias in aliases:
            if alias in header_idx:
                resolved[field_name] = header_idx[alias]
                break
    missing = [f for f in required if f not in resolved]
    if missing:
        raise BusinessDataError(
            f"la hoja de {label} no tiene columna(s) reconocible(s) para {missing}. "
            f"cabeceras encontradas: {sorted(header_idx)}"
        )
    return resolved


def _cell(row: tuple, idx: int | None):
    """Safe positional read: tolerate short rows and unresolved (optional) columns."""
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def normalize_iban(iban: str | None) -> str | None:
    if iban is None:
        return None
    return re.sub(r"\s+", "", str(iban)).upper()


def _to_decimal(value) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        # Excel gives floats; str() gives the shortest exact repr (10325.9 -> "10325.9")
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None


@dataclass
class BusinessData:
    suppliers_by_id: dict[str, Supplier]
    suppliers_by_nif: dict[str, Supplier]
    orders: dict[str, PurchaseOrder]
    rules_version: str
    rules_text: list[str]
    warnings: list[str] = field(default_factory=list)

    def supplier_for_invoice(self, nif: str | None) -> Supplier | None:
        """Rule 1: a real supplier is matched by NIF (invoices carry NIF, not ID)."""
        if not nif:
            return None
        return self.suppliers_by_nif.get(nif.strip())

    def order(self, pedido: str | None) -> PurchaseOrder | None:
        if not pedido:
            return None
        return self.orders.get(pedido.strip())


def load_business_data(xlsx_path: Path = DEFAULT_XLSX) -> BusinessData:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    warnings: list[str] = []

    # --- suppliers (master) ---
    suppliers_by_id: dict[str, Supplier] = {}
    suppliers_by_nif: dict[str, Supplier] = {}
    ws = _find_sheet(wb, SHEET_SUPPLIERS, SUPPLIER_SHEET_ALIASES, warnings, "proveedores")
    scol = _resolve_columns(_header_index(ws), SUPPLIER_COLUMNS, SUPPLIER_REQUIRED, "proveedores")
    for row in ws.iter_rows(min_row=2, values_only=True):
        sid = _cell(row, scol["id"])
        if not sid:
            continue
        name = _cell(row, scol["name"])
        tax_id = _cell(row, scol["tax_id"])
        iban = _cell(row, scol["iban"])
        terms = _cell(row, scol.get("payment_terms"))
        supplier = Supplier(
            id=str(sid).strip(),
            name=str(name).strip() if name else "",
            tax_id=str(tax_id).strip() if tax_id else "",
            iban=normalize_iban(iban) or "",
            payment_terms=str(terms).strip() if terms else None,
        )
        if supplier.id in suppliers_by_id:
            existing = suppliers_by_id[supplier.id]
            if existing.model_dump() != supplier.model_dump():
                warnings.append(f"proveedor {supplier.id} duplicado con datos EN CONFLICTO - se conservo el primero")
            else:
                warnings.append(f"proveedor {supplier.id} duplicado (identico) - de-duplicado")
            continue
        suppliers_by_id[supplier.id] = supplier
        if supplier.tax_id:
            if supplier.tax_id in suppliers_by_nif:
                warnings.append(f"NIF {supplier.tax_id} apunta a mas de un proveedor")
            suppliers_by_nif[supplier.tax_id] = supplier

    # --- purchase orders ---
    orders: dict[str, PurchaseOrder] = {}
    ws = _find_sheet(wb, SHEET_ORDERS, ORDER_SHEET_ALIASES, warnings, "pedidos")
    ocol = _resolve_columns(_header_index(ws), ORDER_COLUMNS, ORDER_REQUIRED, "pedidos")
    for row in ws.iter_rows(min_row=2, values_only=True):
        pedido = _cell(row, ocol["id"])
        if not pedido:
            continue
        raw_amount = _cell(row, ocol["amount"])
        amount = _to_decimal(raw_amount)
        if amount is None:
            warnings.append(f"pedido {pedido} con importe no interpretable {raw_amount!r} - omitido")
            continue
        supplier_id = _cell(row, ocol["supplier_id"])
        tax_id = _cell(row, ocol.get("tax_id"))
        po = PurchaseOrder(
            id=str(pedido).strip(),
            supplier_id=str(supplier_id).strip() if supplier_id else "",
            tax_id=str(tax_id).strip() if tax_id else None,
            amount=amount,
        )
        if po.id in orders:
            warnings.append(f"pedido {po.id} aparece mas de una vez en la hoja de pedidos")
        orders[po.id] = po

    # --- rules text (kept for the pitch / versioning; the logic is coded in rules_engine) ---
    rules_text: list[str] = []
    if SHEET_RULES in wb.sheetnames:
        for row in wb[SHEET_RULES].iter_rows(values_only=True):
            if row and row[0]:
                rules_text.append(str(row[0]).strip())
    rules_version = "norma-v3"

    wb.close()
    return BusinessData(
        suppliers_by_id=suppliers_by_id,
        suppliers_by_nif=suppliers_by_nif,
        orders=orders,
        rules_version=rules_version,
        rules_text=rules_text,
        warnings=warnings,
    )


def _main() -> int:
    biz = load_business_data()
    print(f"suppliers : {len(biz.suppliers_by_id)} (by NIF: {len(biz.suppliers_by_nif)})")
    print(f"orders    : {len(biz.orders)}")
    print(f"rules     : {biz.rules_version} ({len(biz.rules_text)} lines)")
    print(f"warnings  : {biz.warnings}")
    any_id = next(iter(biz.suppliers_by_id))
    print("sample supplier:", biz.suppliers_by_id[any_id].model_dump())
    any_po = next(iter(biz.orders))
    print("sample order   :", biz.orders[any_po].model_dump())
    return 0


if __name__ == "__main__":
    raise SystemExit(_main())
