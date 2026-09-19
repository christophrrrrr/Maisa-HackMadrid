"""Universal document ingestion - "any file in, always a routed decision".

Normalises an arbitrary file into text (+ images + attachments), then runs a
two-pass extraction into the SAME canonical `InvoiceData`:

  Pass A - deterministic regex over the normalised text (free, auditable).
  Pass B - a generic LLM pass over the text/images when A is not confident.

Anything we still cannot read as an invoice comes back `extraction_ok=False`
with a precise triage code, so the deterministic rules engine ESCALATES it with
a reason (out_of_scope / unreadable / unknown_format). The LLM is only ever an
*extraction* fallback - it never makes the payment decision.

Supported here: .txt/.md/.json/.html, .csv, .xlsx, .docx, .eml and .msg.
PDFs and images keep using `invoice_extractor`.
Email attachments that are PDFs/images are run back through the PDF path, so an
invoice hidden inside an email is still decided.
"""
from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from . import invoice_extractor as ax
from . import vision_gemini as vg
from .models import InvoiceData

# reuse the invoice-domain constant so email/text share the "payer, not supplier" rule
CLIENT_CIF = "A58231074"

_TEXT_EXTS = {".txt", ".md", ".json", ".log"}
_HTML_EXTS = {".htm", ".html"}
_CSV_EXTS = {".csv"}
_XLSX_EXTS = {".xlsx"}
_DOCX_EXTS = {".docx"}
_EML_EXTS = {".eml"}
_MSG_EXTS = {".msg"}
_PDF_LIKE = {".pdf"}
_IMAGE_LIKE = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp", ".gif"}

_GENERIC_PROMPT = (
    "You are reading a business document that MAY be a supplier invoice for the payer "
    f"Banco Miralmar (CIF {CLIENT_CIF} - NEVER the supplier). Extract the supplier-invoice "
    "fields if they are present. If the document is clearly NOT a supplier invoice (a contract, "
    "a plain email, a report, a price list...), return null for every field. Copy values exactly "
    "as printed and never infer missing ones. Return issue_date as YYYY-MM-DD, monetary values as "
    "decimal numbers without currency or thousands separators, iva_rate as the percentage number "
    "(for example 21), IBAN without spaces, and the purchase order exactly as printed "
    "(normally PO-2026-NNNN)."
)

# fields that signal "this looks invoice-ish" even when incomplete
_SIGNAL_FIELDS = (
    "purchase_order", "supplier_tax_id", "total",
    "supplier_iban", "base", "iva_amount", "issue_date",
)
_REQUIRED = ("purchase_order", "supplier_tax_id", "total")


@dataclass
class RawDoc:
    text: str = ""
    images: list[bytes] = field(default_factory=list)
    attachments: list[tuple[str, bytes]] = field(default_factory=list)
    source: str = "unknown"
    note: str | None = None


# ------------------------------- adapters ----------------------------------

def _read_bytes(path: Path) -> bytes:
    return path.read_bytes()


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "iso-8859-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _html_to_text(raw: str) -> str:
    try:
        from bs4 import BeautifulSoup
        return BeautifulSoup(raw, "html.parser").get_text("\n")
    except Exception:
        import re
        return re.sub(r"<[^>]+>", " ", raw)


def _csv_to_text(path: Path) -> str:
    import csv
    lines: list[str] = []
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as fh:
        for row in csv.reader(fh):
            lines.append(" | ".join(c.strip() for c in row))
    return "\n".join(lines)


def _xlsx_to_text(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out: list[str] = []
    for ws in wb.worksheets:
        out.append(f"# hoja: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)


def _docx_to_text(path: Path) -> str:
    import docx  # python-docx
    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _eml_to_doc(path: Path) -> RawDoc:
    import email
    from email import policy as email_policy

    msg = email.message_from_bytes(_read_bytes(path), policy=email_policy.default)
    header = []
    for h in ("From", "To", "Subject", "Date"):
        v = msg.get(h)
        if v:
            header.append(f"{h}: {v}")
    text_parts: list[str] = ["\n".join(header)] if header else []
    images: list[bytes] = []
    attachments: list[tuple[str, bytes]] = []

    for part in msg.walk():
        if part.is_multipart():
            continue
        ctype = (part.get_content_type() or "").lower()
        disp = (part.get_content_disposition() or "").lower()
        fname = part.get_filename()
        try:
            payload = part.get_payload(decode=True)
        except Exception:
            payload = None

        if disp == "attachment" or fname:
            if payload:
                attachments.append((fname or "adjunto", payload))
                if ctype.startswith("image/"):
                    images.append(payload)
            continue
        if ctype == "text/plain" and payload:
            text_parts.append(_decode(payload))
        elif ctype == "text/html" and payload:
            text_parts.append(_html_to_text(_decode(payload)))

    return RawDoc(text="\n".join(t for t in text_parts if t),
                  images=images, attachments=attachments, source="email")


def _msg_to_doc(path: Path) -> RawDoc | None:
    try:
        import extract_msg  # optional dependency
    except Exception:
        return None
    m = extract_msg.Message(str(path))
    text = "\n".join(x for x in (m.subject, m.body) if x)
    attachments: list[tuple[str, bytes]] = []
    images: list[bytes] = []
    for att in getattr(m, "attachments", []) or []:
        data = getattr(att, "data", None)
        name = getattr(att, "longFilename", None) or getattr(att, "shortFilename", None) or "adjunto"
        if isinstance(data, bytes):
            attachments.append((name, data))
            if Path(name).suffix.lower() in _IMAGE_LIKE:
                images.append(data)
    return RawDoc(text=text, images=images, attachments=attachments, source="email")


def normalize(path: Path) -> RawDoc | None:
    """Reduce a file to text (+images +attachments). None => no adapter matched."""
    ext = path.suffix.lower()
    try:
        if ext in _EML_EXTS:
            return _eml_to_doc(path)
        if ext in _MSG_EXTS:
            return _msg_to_doc(path)
        if ext in _DOCX_EXTS:
            return RawDoc(text=_docx_to_text(path), source="docx")
        if ext in _XLSX_EXTS:
            return RawDoc(text=_xlsx_to_text(path), source="spreadsheet")
        if ext in _CSV_EXTS:
            return RawDoc(text=_csv_to_text(path), source="spreadsheet")
        if ext in _HTML_EXTS:
            return RawDoc(text=_html_to_text(_decode(_read_bytes(path))), source="html")
        if ext in _TEXT_EXTS:
            return RawDoc(text=_decode(_read_bytes(path)), source="text")
        # unknown extension: best-effort read as text; binary => give up cleanly
        data = _read_bytes(path)
        text = _decode(data)
        printable = sum(c.isprintable() or c.isspace() for c in text[:2000])
        if text and printable / max(len(text[:2000]), 1) > 0.85:
            return RawDoc(text=text, source="raw-text", note="formato desconocido leido como texto")
        return None
    except Exception as exc:
        return RawDoc(source="error", note=f"no se pudo leer: {exc}")


# ---------------------------- extraction core ------------------------------

def _signals(inv: InvoiceData | None) -> int:
    if inv is None:
        return 0
    return sum(1 for f in _SIGNAL_FIELDS if getattr(inv, f, None) not in (None, ""))


def _has_required(inv: InvoiceData | None) -> bool:
    return inv is not None and all(getattr(inv, f, None) not in (None, "") for f in _REQUIRED)


def _try_attachment(name: str, blob: bytes, *, use_vision: bool, model: str,
                    fallback_models: list[str] | None) -> tuple[InvoiceData | None, str, float]:
    """Run a PDF/image email attachment back through the PDF extractor."""
    suffix = Path(name).suffix.lower()
    if suffix not in _PDF_LIKE | _IMAGE_LIKE:
        return None, "", 0.0
    tmp = Path(tempfile.gettempdir()) / f"maisa_att_{abs(hash(name)) % 10_000}{suffix}"
    try:
        tmp.write_bytes(blob)
        rec = ax.extract_pdf(tmp, use_vision=use_vision, model=model, fallback_models=fallback_models)
        return rec.invoice, f"email-attachment:{rec.method}", float(getattr(rec, "cost_usd", 0.0) or 0.0)
    except Exception:
        return None, "", 0.0
    finally:
        try:
            tmp.unlink()
        except OSError:
            pass


def _llm_over_text(text: str, images: list[bytes]) -> tuple[InvoiceData | None, float]:
    """Pass B: generic LLM extraction. Gemini (free) preferred, Gateway fallback."""
    google_key = vg.gemini_available()
    gateway_key = os.getenv("AI_GATEWAY_API_KEY") or os.getenv("VERCEL_OIDC_TOKEN")
    prompt = _GENERIC_PROMPT + "\n\n---DOCUMENT TEXT---\n" + (text or "")[:16000]
    try:
        if google_key:
            raw = vg.call_gemini_json(images, prompt, api_key=google_key)
            parsed = ax.VisionInvoice.model_validate_json(raw)
            inv, _ = ax._invoice_from_vision("", parsed)
            return inv, 0.0
        if gateway_key:
            parsed = _gateway_text(prompt, images, gateway_key)
            inv, _ = ax._invoice_from_vision("", parsed)
            return inv, ax._last_vision_cost
    except Exception:
        return None, 0.0
    return None, 0.0


def _gateway_text(prompt: str, images: list[bytes], gateway_key: str) -> "ax.VisionInvoice":
    import base64
    from openai import OpenAI

    content: list[dict[str, Any]] = [{"type": "text", "text": prompt}]
    for image in images:
        enc = base64.b64encode(image).decode("ascii")
        content.append({"type": "image_url",
                        "image_url": {"url": f"data:image/png;base64,{enc}", "detail": "high"}})
    client = OpenAI(api_key=gateway_key, base_url=ax.GATEWAY_BASE_URL)
    resp = client.chat.completions.create(
        model=ax.DEFAULT_MODEL,
        messages=[{"role": "user", "content": content}],
        response_format={"type": "json_schema", "json_schema": {
            "name": "invoice_extraction", "strict": True,
            "schema": ax.VisionInvoice.model_json_schema()}},
    )
    ax._last_vision_cost = ax._usd_from_usage(ax.DEFAULT_MODEL, getattr(resp, "usage", None))
    txt = resp.choices[0].message.content
    if not txt:
        raise ValueError("gateway returned no content")
    return ax.VisionInvoice.model_validate_json(txt)


def extract_document(
    path: Path,
    *,
    use_vision: bool = True,
    model: str = ax.DEFAULT_MODEL,
    fallback_models: list[str] | None = None,
) -> tuple[InvoiceData, str, float]:
    """Return (InvoiceData, method, cost_usd). Never raises; always routes."""
    raw = normalize(path)
    if raw is None:
        return (InvoiceData(file_id=path.name, extraction_ok=False,
                            extraction_reason="unknown_format",
                            extraction_note=f"sin lector para {path.suffix or 'archivo'}"),
                "unknown_format", 0.0)
    if raw.note and raw.source == "error":
        return (InvoiceData(file_id=path.name, extraction_ok=False,
                            extraction_reason="unreadable", extraction_note=raw.note),
                "unreadable", 0.0)

    total_cost = 0.0

    # 0) email attachments that are invoices win outright
    for name, blob in raw.attachments:
        inv, method, cost = _try_attachment(name, blob, use_vision=use_vision,
                                             model=model, fallback_models=fallback_models)
        total_cost += cost
        if _has_required(inv):
            inv.file_id = path.name
            return inv, method, total_cost

    text = ax._normalise_text(raw.text or "")
    best: InvoiceData | None = None
    best_method = f"text:{raw.source}"

    # Pass A - deterministic regex over the normalised text
    if ax._compact_chars(text) >= ax.MIN_TEXT_CHARS:
        inv_a, *_ = ax._parse_digital(path.name, [text])
        if _has_required(inv_a):
            return inv_a, f"text:{raw.source}", total_cost
        best = inv_a

    # Pass B - generic LLM over text (+ any images)
    if use_vision and (text or raw.images):
        inv_b, cost = _llm_over_text(text, raw.images)
        total_cost += cost
        if inv_b is not None:
            inv_b.file_id = path.name
            if _has_required(inv_b):
                return inv_b, f"llm:{raw.source}", total_cost
            if _signals(inv_b) > _signals(best):
                best, best_method = inv_b, f"llm:{raw.source}"

    # Triage - we could not confidently read an invoice
    if best is None:
        best = InvoiceData(file_id=path.name)
    best.file_id = path.name
    best.extraction_ok = False
    sig = _signals(best)
    if not text and not raw.images:
        best.extraction_reason = "unreadable"
        best.extraction_note = "documento vacio o ilegible"
        best_method = "unreadable"
    elif sig == 0:
        best.extraction_reason = "out_of_scope"
        best.extraction_note = "el documento no parece una factura de proveedor"
        best_method = f"triage:{raw.source}"
    else:
        best.extraction_reason = "incomplete_extraction"
        missing = [f for f in _REQUIRED if getattr(best, f, None) in (None, "")]
        best.extraction_note = "factura incompleta: falta " + ", ".join(missing)
    return best, best_method, total_cost
